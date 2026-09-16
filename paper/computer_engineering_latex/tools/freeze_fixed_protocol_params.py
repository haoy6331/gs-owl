from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


FIXED_VALUES = ("5", "0.08", "0.20", "0.5")

TEXT_REPLACEMENTS = {
    "其中尚未确定的取值暂以统一占位符标记。": (
        "本研究统一固定为(H_m, λ, α, β)=(5,0.08,0.20,0.5)。"
    ),
    "固定参数组合确定并完成同协议测试后": "使用上述固定参数完成同协议测试后",
    "固定参数组合及其": "采用上述固定参数获得的",
    "固定参数组合确定后补入": "采用上述固定参数完成测试后补入",
    "固定参数组合确定后再补入": "完成上述固定配置的评测后再补入",
    "固定参数组合、六个唯一Mask": "固定参数组合已确定，六个唯一Mask",
}


def replace_in_runs(paragraph, old: str, new: str) -> int:
    count = 0
    for run in paragraph.runs:
        if old in run.text:
            occurrences = run.text.count(old)
            run.text = run.text.replace(old, new)
            count += occurrences
    return count


def set_cell_value(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = value
        first.font.highlight_color = None
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    args = parser.parse_args()

    document = Document(args.input_docx)

    if len(document.tables) < 2 or len(document.tables[1].rows) < 5:
        raise RuntimeError("The fixed-parameter table was not found at the expected location.")

    parameter_table = document.tables[1]
    for row_index, value in enumerate(FIXED_VALUES, start=1):
        set_cell_value(parameter_table.cell(row_index, 1), value)

    applied = {old: 0 for old in TEXT_REPLACEMENTS}
    for paragraph in document.paragraphs:
        for old, new in TEXT_REPLACEMENTS.items():
            applied[old] += replace_in_runs(paragraph, old, new)

    missing = [old for old, count in applied.items() if count == 0]
    if missing:
        raise RuntimeError(f"Expected text was not found: {missing}")

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)

    print(f"Saved: {args.output_docx}")
    for old, count in applied.items():
        print(f"Replaced {count} occurrence(s): {old}")


if __name__ == "__main__":
    main()
